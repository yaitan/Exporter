import sys
import requests
import openpyxl as xl
import datetime
import os
from openpyxl.worksheet.table import Table

LOGIN_REDIRECT_URL = "YOUR_LOGIN_REDIRECT_URL"

LOGIN_URL = "YOUR_INSTAGRAM_BUSINESS_LOGIN_URL"

CLIENT_ID = "YOUR_CLIENT_ID"

API = "https://graph.instagram.com/v23.0"

MEDIA_FIELDS = ['id',
                'caption',
                'media_product_type',
                'media_type',
                'permalink',
                'timestamp',
                'like_count',
                'comments_count',
                'is_shared_to_feed'
                ]

COMMENT_FIELDS = ['text',
                  'id',
                  'like_count',
                  'replies',
                  'timestamp',
                  'parent_id',
                  'hidden',
                  ]

STANDARDIZED_COMMENT_FIELDS = {'id': None,
                               'post_id': 0,
                               'text': '',
                               'like_count': 0,
                               'replies': {'data': []},
                               'timestamp': None,
                               'parent_id': None,
                               'hidden': 'UNKNOWN'}


def reformat_media_product_type(json):
    """
    combines fields media_product_type and is_shared_to_feed into one field
    :param json:
    :return:
    """
    product = json['media_product_type']
    if product == "REELS":
        feed = json['is_shared_to_feed']
        json['media_product_type'] = 'Reel in feed' if feed else 'Reel'
    else:
        json['media_product_type'] = 'Post'
    return json


def create_xl(name):
    """
    creates excel workbook with first sheet named Posts and saves the workbook
    :param name: filename of excel file
    :return: workbook object
    """
    wb = xl.Workbook()
    sheet = wb.active
    sheet.title = "Posts"
    sheet.append(MEDIA_FIELDS[:-1])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20
    wb.save(name)
    return wb


def add_all_media(media_list):
    """
    adds all media from media_list to the Excel workbook and saves the workbook
    :param media_list: list of media_ids
    """
    count = 1
    sheet = workbook.active
    for media in media_list:
        data = get_media_info(media['id'])
        data['timestamp'] = reformat_time(data['timestamp'])
        sheet.append(list(data.values())[:-1])

        print('\r', end='')
        print(f'{count}/{len(media_list)}', end='')
        count += 1
    print("\nAdded media")
    table = Table(displayName="Posts", ref=f"A1:H{len(media_list) + 1}")
    sheet.add_table(table)
    workbook.save(filename)


def reformat_time(time):
    """
    reformats time from string to excel format
    :param time: string in format %Y-%m-%dT%H:%M:%S%z
    :return: string in format %Y-%m-%d %H:%M:%S
    """
    dt = datetime.datetime.strptime(time, '%Y-%m-%dT%H:%M:%S%z')
    excel_format = dt.strftime('%Y-%m-%d %H:%M:%S')
    return excel_format


def create_comments_sheet():
    """
    creates sheet named Comments in workbook and saves the workbook
    :return: sheet object
    """
    sheet = workbook.create_sheet("Comments")
    sheet.append(list(STANDARDIZED_COMMENT_FIELDS.keys()))
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20
    workbook.save(filename)
    return sheet


def standarize_comment(comment, media_id):
    """
    ensures that all comments have the same structure, and changes replies to number of replies
    :param comment: Json formatted comment
    :param media_id: id of media post comment belongs to
    :return: standardized version of the comment
    """
    stand_comment = {k: comment.get(k, v) for k, v in STANDARDIZED_COMMENT_FIELDS.items()}
    stand_comment['text'] = "'" + stand_comment['text']
    stand_comment['post_id'] = media_id
    stand_comment['replies'] = len(stand_comment['replies']['data'])
    stand_comment['timestamp'] = reformat_time(stand_comment['timestamp'])
    return stand_comment


def add_comments(media_id, sheet):
    """
    adds comments from the media to the comments sheet and saves the workbook
    :param media_id: id of media post to get comments from
    :param sheet: sheet object to add comments to
    :return: number of comments added to sheet
    """
    comments = get_comments(media_id)
    for comment in comments:
        comment = standarize_comment(comment, media_id)
        sheet.append(list(comment.values()))
    workbook.save(filename)
    return len(comments)


def get_media_ids():
    """
    gets all media ids from the account from Instagram API
    :return: list of media ids
    """
    data = []
    url = f"{API}/me/media?{ACCESS}&limit=100"
    response = requests.get(url)
    json = response.json()
    data += json['data']
    print(f'fetched {len(data)} posts')

    while 'next' in json['paging']:
        response = requests.get(json['paging']['next'])
        json = response.json()
        data += json['data']
        print(f'fetched {len(data)} posts')
    return data


def get_comments(media_id):
    """
    gets all comments for a specific media id from Instagram API
    :param media_id: media id to get comments for
    :return: list of comments in json format, each comment is a dictionary with keys from comment_fields
    """
    data = []
    fields = ','.join(COMMENT_FIELDS)
    url = f"{API}/{media_id}/comments?fields={fields}&{ACCESS}&limit=100"
    response = requests.get(url)
    json = response.json()
    data += json.get('data', 0)
    while 'next' in json.get('paging', []):
        response = requests.get(json['paging']['next'])
        json = response.json()
        data += json['data']
    return data


def get_media_info(media_id):
    """
    gets all media info for a specific media id from Instagram API
    :param media_id:
    :return: json with keys from media_fields
    """
    fields = ",".join(MEDIA_FIELDS)
    url = f"{API}/{media_id}/?fields={fields}&{ACCESS}"
    response = requests.get(url)
    json = reformat_media_product_type(response.json())
    return json


def refresh_token(access):
    """
    Refreshes access token. If it fails, generates a new one and returns it.
    :param access: Access token as a formatted string
    :return: new access token as a formatted string
    """
    url = f"{API}/refresh_access_token?grant_type=ig_refresh_token&{access}"
    response = requests.get(url)
    json = response.json()
    if 'access_token' in json:
        with open(access_token_path, 'w') as access_file:
            access_file.write(json['access_token'])
        return f"access_token={json['access_token']}"
    else:
        return generate_access_token()


def generate_temp_token(secret):
    """
    generates a short lived access token using instagram login.
    :param secret: app secret
    :return: short lived access token
    """
    print("Access token expired. Please log in again. Copy the link into your browser,"
          " then login from the desired account")
    print(LOGIN_URL)
    response_url = input("Enter URL of the redirected page (it should be a 404 page)\n")
    code = response_url.split("=")[1].split("#")[0]

    url = "https://api.instagram.com/oauth/access_token"
    data = {
        'client_id': CLIENT_ID,
        'client_secret': secret,
        'grant_type': 'authorization_code',
        'redirect_uri': LOGIN_REDIRECT_URL,
        'code': code
    }
    response = requests.post(url, data=data)
    json = response.json()
    if 'access_token' in json:
        return json['access_token']
    else:
        print("Failed to generate token. Please try again.")
        return generate_temp_token(secret)


def exchange_access_token(temp_token, secret):
    """
    Exchanges short lived access token for long lived access token..
    :param temp_token: short lived access token
    :param secret: app secret
    :return: long lived access token as a formatted string
    """
    url = f"https://graph.instagram.com/access_token"
    params = {
        'grant_type': 'ig_exchange_token',
        'client_secret': secret,
        'access_token': temp_token
    }
    json = requests.get(url, params=params).json()
    if 'access_token' in json:
        with open(access_token_path, 'w') as access_file:
            access_file.write(json['access_token'])
        print("Access token generated successfully")
        return f"access_token={json['access_token']}"
    else:
        print("Failed to exchange access token. Please try again.")
        return generate_temp_token(secret)


def generate_access_token():
    """
    Generates a new access token through instagram login.
    :return: new access token as a formatted string
    """
    with open(secret_path, 'r') as secret_file:
        secret = secret_file.read()
    temp_token = generate_temp_token(secret)
    return exchange_access_token(temp_token, secret)


def create_file_paths():
    """
    creates filepaths for access token and secret file for use in exe
    :return: useable filepaths
    """
    access_file_path = os.path.join(os.getcwd(), "access_token.txt")

    if getattr(sys, 'frozen', False):
        # If the application is run as a bundle
        base_path = sys._MEIPASS
    else:
        # If run normally
        base_path = os.path.abspath(".")

    secret_path_file = os.path.join(base_path, 'app secret')
    return access_file_path, secret_path_file


if __name__ == '__main__':
    # get access token
    access_token_path, secret_path = create_file_paths()
    with open(access_token_path, 'r') as file:
        access_token = file.read()
    ACCESS = f"access_token={access_token}"
    ACCESS = refresh_token(ACCESS)
    # create xlsx file
    date = datetime.datetime.now().strftime('%d_%m_%y %H%M')
    filename = f"Snapshot {date}.xlsx"
    workbook = create_xl(filename)
    create_comments_sheet()
    # get posts
    medias = get_media_ids()
    add_all_media(medias)
    # get comments
    comments_read = 0
    for i, m in enumerate(medias):
        comments_read += add_comments(m['id'], workbook['Comments'])
        print('\r', end='')
        print(f'Fetched {comments_read} from {i + 1} posts', end='')
    tab = Table(displayName="Comments", ref=f"A1:H{comments_read + 1}")
    workbook['Comments'].add_table(tab)
    workbook.save(filename)

    print('\rFinished')
